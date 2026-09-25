#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
pipeline.py — Lean one-pass Telugu SBV/SRT -> English review workbook.

Reusable core shared by the CLI (cli.py) and the web app (app.py).
Uses Azure OpenAI. Configure via environment variables (see .env.example):

    AZURE_OPENAI_ENDPOINT=https://your-resource.openai.azure.com
    AZURE_OPENAI_API_KEY=...
    AZURE_OPENAI_API_VERSION=2024-10-21
    AZURE_OPENAI_DEPLOYMENT=your-deployment-name

Produces:
    - master review workbook (.xlsx)
    - discourse brief (.md)
    - raw translation (.json)

After human review, the same rows produce the final English .srt / .sbv, and an
optional AI review pass (MQM-style) flags likely errors for the human to decide on.
"""

from __future__ import annotations

import json
import math
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Tuple

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from openai import AzureOpenAI
except Exception:  # pragma: no cover - import guard
    AzureOpenAI = None  # type: ignore

DEFAULT_CHUNK_SIZE = 40
DEFAULT_OVERLAP = 3
DEFAULT_SKILL_FILE = "SKILL_v4_1_lean_subtitle_review_pipeline.md"
CHUNK_ATTEMPTS = 3              # re-ask the model when a chunk comes back malformed
REVIEW_CHUNK_SIZE = 10          # reasoning time grows much faster than chunk size: a
                                # 10-cue review takes ~40 s, a 40-cue one 13+ min
LLM_TIMEOUT_SEC = 600           # a stuck request fails (and is retried) instead of hanging
MAX_LINE_CHARS = 42             # Netflix Timed Text Style Guide: 42 chars/line, 2 lines
MAX_CPS = 20                    # Netflix English: up to 20 characters per second (adult)

WORKBOOK_COLUMNS = [
    "Cue Number",
    "Timecode",
    "Telugu Cue",
    "AI Generated English Cue",
    "Human Review Correction",
]

# progress(message) -> None. Optional hook so callers (web UI) can show status.
ProgressFn = Callable[[str], None]


def _noop(_msg: str) -> None:
    pass


@dataclass
class Cue:
    number: int
    start: str
    end: str
    telugu: str

    @property
    def timecode(self) -> str:
        return f"{self.start} --> {self.end}"

    @property
    def duration(self) -> float:
        return ts_to_seconds(self.end) - ts_to_seconds(self.start)


# --------------------------------------------------------------------------- #
# Timestamps
# --------------------------------------------------------------------------- #
def ts_to_seconds(ts: str) -> float:
    """Parse 'HH:MM:SS,mmm', 'H:MM:SS.mmm' or 'MM:SS.mmm' into seconds."""
    parts = ts.strip().replace(",", ".").split(":")
    if not 2 <= len(parts) <= 3:
        raise ValueError(f"Bad timestamp: {ts!r}")
    hours = int(parts[0]) if len(parts) == 3 else 0
    minutes, seconds = int(parts[-2]), float(parts[-1])
    if minutes >= 60 or seconds >= 60 or min(hours, minutes, seconds) < 0:
        raise ValueError(f"Bad timestamp: {ts!r}")
    return hours * 3600 + minutes * 60 + seconds


def _split_ms(seconds: float) -> Tuple[int, int, int, int]:
    ms = int(round(max(0.0, seconds) * 1000))
    h, ms = divmod(ms, 3_600_000)
    m, ms = divmod(ms, 60_000)
    s, ms = divmod(ms, 1000)
    return h, m, s, ms


def seconds_to_srt_ts(seconds: float) -> str:
    h, m, s, ms = _split_ms(seconds)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def seconds_to_sbv_ts(seconds: float) -> str:
    h, m, s, ms = _split_ms(seconds)
    return f"{h}:{m:02d}:{s:02d}.{ms:03d}"


def _ts_sbv_to_srt(ts: str) -> str:
    return seconds_to_srt_ts(ts_to_seconds(ts))


def _normalize_srt_ts(ts: str) -> str:
    return seconds_to_srt_ts(ts_to_seconds(ts))


# --------------------------------------------------------------------------- #
# Parsing
# --------------------------------------------------------------------------- #
def parse_source_text(raw: str) -> List[Cue]:
    """Parse SBV or SRT subtitle text into ordered cues."""
    raw = raw.lstrip("﻿")
    # Whitespace-only separator lines are common in hand-edited files.
    blocks = [b for b in re.split(r"\r?\n[ \t]*\r?\n", raw) if b.strip()]
    cues: List[Cue] = []

    for idx, block in enumerate(blocks, 1):
        lines = [ln.rstrip() for ln in block.splitlines() if ln.strip() != ""]
        if not lines:
            continue

        if lines[0].strip().isdigit() and len(lines) > 1:
            timing_line = lines[1].strip()
            text_lines = lines[2:]
        else:
            timing_line = lines[0].strip()
            text_lines = lines[1:]

        if "-->" in timing_line:
            start_raw, end_raw = [x.strip() for x in timing_line.split("-->", 1)]
            # SRT allows position settings after the end time ("X1:... Y1:...").
            start, end = _normalize_srt_ts(start_raw), _normalize_srt_ts(end_raw.split()[0])
        else:
            if "," not in timing_line:
                raise ValueError(f"Could not parse timing at block {idx}: {timing_line!r}")
            start_raw, end_raw = timing_line.split(",", 1)
            start = _ts_sbv_to_srt(start_raw)
            end = _ts_sbv_to_srt(end_raw)

        cues.append(Cue(number=idx, start=start, end=end, telugu="\n".join(text_lines).strip()))

    return [Cue(i, c.start, c.end, c.telugu) for i, c in enumerate(cues, 1)]


def parse_source(path: Path) -> List[Cue]:
    return parse_source_text(Path(path).read_text(encoding="utf-8-sig"))


def safe_stem(name: str) -> str:
    stem = Path(name).stem
    return re.sub(r"[^A-Za-z0-9_-]+", "_", stem).strip("_") or "subtitle"


# --------------------------------------------------------------------------- #
# Writing subtitles
# --------------------------------------------------------------------------- #
def wrap_subtitle(text: str, max_chars: int = MAX_LINE_CHARS) -> str:
    """Keep author line breaks; otherwise split long text into two balanced lines."""
    text = text.strip()
    if "\n" in text:
        return "\n".join(" ".join(line.split()) for line in text.splitlines() if line.strip())
    text = " ".join(text.split())
    if len(_strip_markup(text)) <= max_chars or " " not in text:
        return text
    mid = len(text) / 2
    best = min((i for i, ch in enumerate(text) if ch == " "), key=lambda i: abs(i - mid))
    return text[:best] + "\n" + text[best + 1:]


def _strip_markup(text: str) -> str:
    return text.replace("*", "")


def _render_italics(text: str, fmt: str) -> str:
    """The workbook marks Sanskrit terms with Markdown *asterisks*; SRT players render
    <i>, while YouTube SBV has no styling, so the markers are dropped there.

    A span may cross the line break added by wrapping, so tags are closed at the end
    of each line and reopened on the next (the usual SRT convention).
    """
    if fmt != "srt" or text.count("*") % 2:  # unbalanced markers: don't guess
        return text.replace("*", "")
    lines, italic = [], False
    for line in text.split("\n"):
        out = "<i>" if italic else ""
        for i, part in enumerate(line.split("*")):
            if i:
                italic = not italic
                out += "<i>" if italic else "</i>"
            out += part
        if italic:
            out += "</i>"
        lines.append(out.replace("<i></i>", ""))
    return "\n".join(lines)


def format_subtitles(entries: Iterable[Tuple[str, str, str]], fmt: str) -> str:
    """entries: (start, end, text) with SRT-style or SBV-style timestamps."""
    if fmt not in ("srt", "sbv"):
        raise ValueError(f"Unknown subtitle format: {fmt}")
    blocks = []
    for i, (start, end, text) in enumerate(entries, 1):
        s, e = ts_to_seconds(start), ts_to_seconds(end)
        body = _render_italics(wrap_subtitle(text), fmt)
        if fmt == "srt":
            blocks.append(f"{i}\n{seconds_to_srt_ts(s)} --> {seconds_to_srt_ts(e)}\n{body}")
        else:
            blocks.append(f"{seconds_to_sbv_ts(s)},{seconds_to_sbv_ts(e)}\n{body}")
    return "\n\n".join(blocks) + "\n"


def load_skill_text(skill_path: Optional[Path]) -> str:
    if skill_path and Path(skill_path).exists():
        return Path(skill_path).read_text(encoding="utf-8")
    default = Path(__file__).parent / DEFAULT_SKILL_FILE
    if default.exists():
        return default.read_text(encoding="utf-8")
    here = Path.cwd() / "SKILL.md"
    if here.exists():
        return here.read_text(encoding="utf-8")
    return "Follow the subtitle translation skill strictly."


def cues_to_text(cues: Iterable[Cue]) -> str:
    return "\n".join(
        f"{c.number}. [{c.timecode}] {' / '.join(c.telugu.splitlines())}"
        for c in cues
    )


def build_windows(cues: List[Cue], chunk_size: int) -> List[Tuple[int, int]]:
    windows = []
    start = 1
    while start <= len(cues):
        end = min(len(cues), start + chunk_size - 1)
        windows.append((start, end))
        start = end + 1
    return windows


# --------------------------------------------------------------------------- #
# Azure OpenAI client
# --------------------------------------------------------------------------- #
def get_client() -> Any:
    if AzureOpenAI is None:
        raise RuntimeError("openai package is not installed. Run: pip install openai")

    load_dotenv()
    load_dotenv(".env.local")

    endpoint = os.getenv("AZURE_OPENAI_ENDPOINT")
    api_key = os.getenv("AZURE_OPENAI_API_KEY")
    api_version = os.getenv("AZURE_OPENAI_API_VERSION") or "2024-10-21"

    if not endpoint or not api_key:
        raise RuntimeError(
            "Missing Azure OpenAI config. Set AZURE_OPENAI_ENDPOINT and "
            "AZURE_OPENAI_API_KEY (see .env.example)."
        )

    return AzureOpenAI(
        azure_endpoint=endpoint,
        api_key=api_key,
        api_version=api_version,
        timeout=float(os.getenv("LLM_TIMEOUT_SEC", LLM_TIMEOUT_SEC)),
    )


def resolve_deployment(model: Optional[str] = None) -> str:
    deployment = model or os.getenv("AZURE_OPENAI_DEPLOYMENT")
    if not deployment:
        raise RuntimeError(
            "No model/deployment set. Provide --model or set AZURE_OPENAI_DEPLOYMENT."
        )
    return deployment


def response_text(resp: Any) -> str:
    try:
        if hasattr(resp, "choices") and resp.choices:
            return resp.choices[0].message.content or ""
    except Exception:
        pass
    if hasattr(resp, "output_text") and resp.output_text:
        return resp.output_text
    return str(resp)


def call_model(client: Any, model: str, system: str, user: str) -> str:
    # Temperature intentionally omitted: some Azure deployments only accept the default.
    last_err: Optional[Exception] = None
    for attempt in range(3):
        try:
            resp = client.chat.completions.create(
                model=model,
                messages=[
                    {"role": "system", "content": system},
                    {"role": "user", "content": user},
                ],
            )
            return response_text(resp).strip()
        except Exception as err:  # retry with backoff
            if getattr(err, "status_code", None) == 404:  # config error: retrying won't help
                raise RuntimeError(
                    f"Azure OpenAI returned 404 for deployment {model!r}. Check AZURE_OPENAI_DEPLOYMENT "
                    "and AZURE_OPENAI_API_VERSION (an API version such as 2024-12-01-preview, "
                    "not a model version date)."
                ) from err
            last_err = err
            if attempt == 2:
                raise
            time.sleep(2 * (attempt + 1))
    raise RuntimeError(f"Model call failed: {last_err}")


def extract_json_array(text: str) -> List[Dict[str, Any]]:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)

    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        m = re.search(r"\[[\s\S]*\]", text)
        if not m:
            raise ValueError(f"Model did not return JSON array. Output starts:\n{text[:500]}")
        data = json.loads(m.group(0))

    if not isinstance(data, list):
        raise ValueError("Model JSON output was not a list.")
    return data


def _with_retries(fn: Callable[[], Any], label: str, progress: ProgressFn) -> Any:
    """Re-ask the model when its output is malformed (bad JSON, wrong cue numbers)."""
    for attempt in range(1, CHUNK_ATTEMPTS + 1):
        try:
            return fn()
        except ValueError as err:
            if attempt == CHUNK_ATTEMPTS:
                raise
            progress(f"{label}: malformed model output ({err}); retrying ({attempt}/{CHUNK_ATTEMPTS - 1})...")


def _run_windows(
    windows: List[Tuple[int, int]],
    work: Callable[[int, int], List[Dict[str, Any]]],
    label: str,
    progress: ProgressFn,
) -> Dict[int, Dict[str, Any]]:
    """Run independent cue windows concurrently; each only needs the brief + source."""
    items: Dict[int, Dict[str, Any]] = {}
    workers = max(1, int(os.getenv("LLM_WORKERS", "4")))
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = {
            pool.submit(_with_retries, lambda s=s, e=e: work(s, e), f"{label} {s}-{e}", progress): (s, e)
            for s, e in windows
        }
        for done, fut in enumerate(as_completed(futures), 1):
            s, e = futures[fut]
            for item in fut.result():
                items[int(item["cue_number"])] = item
            progress(f"{label}: cues {s}-{e} done ({done}/{len(windows)} chunks).")
    return items


# --------------------------------------------------------------------------- #
# Translation
# --------------------------------------------------------------------------- #
def make_discourse_brief(client: Any, model: str, skill_text: str, cues: List[Cue]) -> str:
    system = "Create concise context briefs for Telugu-to-English subtitle translation."
    user = f"""
SKILL:
{skill_text}

Create a concise discourse brief for the Telugu episode below.

Include only what helps translation:
- teaching flow
- questions/doubts/prayers/reflections
- examples/analogies
- Sanskrit terms or quotations
- speaker perspective/voice notes
- opening and closing/sign-off cues
- likely cue-boundary risks

Keep it concise. Do not translate cue-by-cue.

TELUGU CUES:
{cues_to_text(cues)}
"""
    return call_model(client, model, system, user)


def translate_chunk(
    client: Any,
    model: str,
    skill_text: str,
    discourse_brief: str,
    cues: List[Cue],
    start: int,
    end: int,
    overlap: int,
) -> List[Dict[str, Any]]:
    before = cues[max(0, start - 1 - overlap): start - 1]
    core = cues[start - 1: end]
    after = cues[end: min(len(cues), end + overlap)]

    system = "You are a faithful Telugu-to-English subtitle translator. Return only valid JSON."
    user = f"""
SKILL:
{skill_text}

DISCOURSE BRIEF:
{discourse_brief}

CONTEXT BEFORE:
{cues_to_text(before) if before else "(none)"}

CORE CUES TO TRANSLATE:
{cues_to_text(core)}

CONTEXT AFTER:
{cues_to_text(after) if after else "(none)"}

TASK:
Translate ONLY the core cues.

Return a JSON array with exactly one object per core cue.

Required object shape:
{{
  "cue_number": 1,
  "english": "English translation here"
}}

Rules:
- Follow SKILL.md strictly.
- Use the Telugu cue as source of truth.
- Translate every core cue; do not leave English blank.
- Do not skip, merge, absorb, reorder, or renumber cues.
- Use context only to avoid drift; do not translate context cues.
- Preserve meaning, teaching intent, examples, questions, Sanskrit terms, and tone.
- Use plain natural spoken English, not literary or dramatic prose.
- Do not intensify ordinary Telugu wording unless Telugu clearly requires it.
- Preserve speaker perspective. If Telugu voices a person's own question, prayer, doubt, or reflection directly, preserve that perspective in English.
- Do not introduce quoted inner speech unless Telugu clearly presents direct speech or quoted thought.
- Use required colon transliteration for Sanskrit names, terms, salutations, and verses wherever they appear.
- Use Markdown asterisks for Sanskrit technical terms and quotations where appropriate.
- Preserve closing/sign-off phrases such as *Jai Sri:manna:ra:yana*.
- Return cue numbers only from {start} to {end}.
"""
    data = extract_json_array(call_model(client, model, system, user))
    expected = list(range(start, end + 1))
    got = [int(x.get("cue_number", -1)) for x in data]
    if got != expected:
        raise ValueError(f"Chunk {start}-{end} cue mismatch. Got {got}, expected {expected}")
    blank = [x["cue_number"] for x in data if not str(x.get("english") or "").strip()]
    if blank:
        raise ValueError(f"blank English for cues {blank}")
    return data


def build_rows(cues: List[Cue], all_items: Dict[int, Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Review rows: the unit the workbook, the in-app editor and the exports share."""
    rows = []
    for cue in cues:
        english = str(all_items[cue.number].get("english", "") or "").strip()
        if not english:
            raise RuntimeError(
                f"AI returned blank English for cue {cue.number}; automated blanks are not allowed."
            )
        rows.append({
            "cue": cue.number,
            "start": cue.start,
            "end": cue.end,
            "telugu": cue.telugu,
            "english": english,
            "correction": "",
            "review": None,
        })
    return rows


def translate_cues(
    cues: List[Cue],
    *,
    skill_text: str,
    model: Optional[str] = None,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    overlap: int = DEFAULT_OVERLAP,
    progress: ProgressFn = _noop,
) -> Dict[str, Any]:
    """Run the full translation pass over parsed cues.

    Returns a dict with keys: brief, rows, all_items.
    """
    if not cues:
        raise RuntimeError("No cues parsed.")

    client = get_client()
    deployment = resolve_deployment(model)

    progress(f"Parsed {len(cues)} cues. Building discourse brief...")
    brief = make_discourse_brief(client, deployment, skill_text, cues)

    windows = build_windows(cues, chunk_size)
    progress(f"Translating {len(windows)} chunks...")
    all_items = _run_windows(
        windows,
        lambda s, e: translate_chunk(client, deployment, skill_text, brief, cues, s, e, overlap),
        "Translation",
        progress,
    )

    missing = [c.number for c in cues if c.number not in all_items]
    if missing:
        raise RuntimeError(f"Missing translated cues: {missing}")

    progress("Assembling review workbook...")
    rows = build_rows(cues, all_items)
    all_items = {k: all_items[k] for k in sorted(all_items)}
    return {"brief": brief, "rows": rows, "all_items": all_items}


# --------------------------------------------------------------------------- #
# AI review (optional second pass)
# --------------------------------------------------------------------------- #
# MQM-style error annotation (themqm.org typology; GEMBA-MQM, WMT 2023). The model
# only flags and suggests; the human reviewer decides (ISO 18587 full post-editing).
REVIEW_CATEGORIES = ("accuracy", "terminology", "linguistic", "style", "readability")
REVIEW_SEVERITIES = ("minor", "major", "critical")


def current_english(row: Dict[str, Any]) -> str:
    return (row.get("correction") or "").strip() or row["english"]


def _review_lines(rows: List[Dict[str, Any]]) -> str:
    out = []
    for r in rows:
        english = current_english(r)
        seconds = max(0.1, ts_to_seconds(r["end"]) - ts_to_seconds(r["start"]))
        cps = len(_strip_markup(english)) / seconds
        telugu = " / ".join(r["telugu"].splitlines())
        out.append(f"{r['cue']}. [{seconds:.1f}s, {cps:.0f} chars/s] TE: {telugu}\n    EN: {english}")
    return "\n".join(out)


def review_chunk(
    client: Any,
    model: str,
    skill_text: str,
    brief: str,
    rows: List[Dict[str, Any]],
    start: int,
    end: int,
) -> List[Dict[str, Any]]:
    before = rows[max(0, start - 3): start - 1]
    core = rows[start - 1: end]
    after = rows[end: end + 2]

    system = (
        "You are a meticulous Telugu-to-English subtitle reviewer using the MQM error "
        "typology. You flag real errors only. Return only valid JSON."
    )
    user = f"""
TRANSLATION STANDARD (the English must follow this):
{skill_text}

DISCOURSE BRIEF:
{brief or "(none)"}

CONTEXT BEFORE (do not review):
{_review_lines(before) if before else "(none)"}

CUES TO REVIEW:
{_review_lines(core)}

CONTEXT AFTER (do not review):
{_review_lines(after) if after else "(none)"}

TASK:
Check each EN line against its Telugu (TE) source. Report an issue only in these categories:
- accuracy: mistranslation, omission, addition, untranslated text, meaning moved into the wrong cue
- terminology: Sanskrit names/terms not in the required colon transliteration, inconsistent terms
- linguistic: grammar, spelling, punctuation
- style: literary/dramatic/sermon-like English, intensified wording, changed speaker perspective, question turned into a statement
- readability: clearly too long to read in the cue time (well over {MAX_CPS} chars/s) and can be shortened without losing meaning

Severity:
- critical: misleads the viewer or distorts the teaching
- major: meaning noticeably wrong, missing or added
- minor: small imperfection; meaning intact

Do not flag acceptable English just because you would phrase it differently. Most cues should be "ok".
A cue that is a fragment of a sentence continuing in the next cue is normal, not an error.

Return a JSON array with exactly one object per cue to review, cue numbers {start} to {end} in order:
{{"cue_number": 1, "verdict": "ok"}}
or
{{"cue_number": 1, "verdict": "issue", "severity": "minor|major|critical",
  "category": "accuracy|terminology|linguistic|style|readability",
  "issue": "one short sentence explaining the problem",
  "suggestion": "the full corrected English for this cue"}}
"""
    data = extract_json_array(call_model(client, model, system, user))
    expected = list(range(start, end + 1))
    got = [int(x.get("cue_number", -1)) for x in data]
    if got != expected:
        raise ValueError(f"Review chunk {start}-{end} cue mismatch. Got {got}, expected {expected}")
    return data


def _normalize_finding(item: Dict[str, Any], reviewed: str) -> Dict[str, Any]:
    if str(item.get("verdict", "ok")).lower() != "issue":
        return {"verdict": "ok", "reviewed": reviewed}
    severity = str(item.get("severity", "minor")).lower()
    category = str(item.get("category", "accuracy")).lower()
    return {
        "verdict": "issue",
        "severity": severity if severity in REVIEW_SEVERITIES else "minor",
        "category": category if category in REVIEW_CATEGORIES else "accuracy",
        "issue": str(item.get("issue") or "").strip(),
        "suggestion": str(item.get("suggestion") or "").strip(),
        "reviewed": reviewed,
    }


def review_rows(
    rows: List[Dict[str, Any]],
    *,
    skill_text: str,
    brief: str,
    model: Optional[str] = None,
    chunk_size: int = REVIEW_CHUNK_SIZE,
    progress: ProgressFn = _noop,
) -> List[Dict[str, Any]]:
    """Attach an MQM-style finding to every row (row["review"]). Returns the rows."""
    client = get_client()
    deployment = resolve_deployment(model)
    windows = build_windows(rows, chunk_size)  # rows are numbered 1..N like cues
    progress(f"AI review of {len(rows)} cues in {len(windows)} chunks...")
    items = _run_windows(
        windows,
        lambda s, e: review_chunk(client, deployment, skill_text, brief, rows, s, e),
        "AI review",
        progress,
    )
    for r in rows:
        r["review"] = _normalize_finding(items.get(r["cue"], {}), current_english(r))
    counts = {sev: sum(1 for r in rows if (r["review"] or {}).get("severity") == sev) for sev in REVIEW_SEVERITIES}
    progress(
        "AI review done: "
        + ", ".join(f"{n} {sev}" for sev, n in counts.items())
        + f" issue(s) across {len(rows)} cues."
    )
    return rows


# --------------------------------------------------------------------------- #
# Workbook
# --------------------------------------------------------------------------- #
# Excel sizes rows for the cell's font. Calibri has no Telugu glyphs, so Excel fell
# back to a taller Indic font inside rows measured for Calibri, clipping conjuncts
# and cutting off wrapped lines. Telugu cells now use a Telugu-capable font and every
# row gets an explicit height, so viewers that do not auto-fit rows (Excel web and
# mobile, Google Sheets imports) also show the full text.
TELUGU_FONT = "Nirmala UI"
REVIEW_WIDTHS = {"A": 12, "B": 30, "C": 64, "D": 66, "E": 48}
# (points per wrapped line, characters per line at the column width) — calibrated in Excel.
LATIN_LINE_PT, LATIN_CHARS_PER_WIDTH = 15.0, 1.1
TELUGU_LINE_PT, TELUGU_CHARS_PER_WIDTH = 17.5, 1.0


def _clean_cell(value: Any) -> Any:
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _wrapped_lines(text: str, chars_per_line: float) -> int:
    return sum(max(1, math.ceil(len(line) / chars_per_line)) for line in str(text).split("\n"))


def _write_sheet(ws: Any, headers: List[str], rows: List[List[Any]], widths: Dict[str, int],
                 telugu_cols: Iterable[str] = ()) -> None:
    telugu_cols = set(telugu_cols)
    ws.append(headers)
    for values in rows:
        ws.append([_clean_cell(v) for v in values])

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(min_row=2):
        line_heights = []
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.data_type == "f":  # text like "=..." must stay text, not a formula
                cell.data_type = "s"
            if cell.value in (None, ""):
                continue
            width = widths.get(cell.column_letter, 10)
            if cell.column_letter in telugu_cols:
                cell.font = Font(name=TELUGU_FONT, size=11)
                lines = _wrapped_lines(cell.value, width * TELUGU_CHARS_PER_WIDTH)
                line_heights.append(lines * TELUGU_LINE_PT)
            else:
                lines = _wrapped_lines(cell.value, width * LATIN_CHARS_PER_WIDTH)
                line_heights.append(lines * LATIN_LINE_PT)
        ws.row_dimensions[row[0].row].height = max(line_heights or [LATIN_LINE_PT]) + 4

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"


def create_workbook(rows: List[Dict[str, Any]], out_path: Path) -> None:
    """Review sheet (the five required columns) + an AI Review sheet when one was run."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    _write_sheet(
        ws,
        WORKBOOK_COLUMNS,
        [[r["cue"], f"{r['start']} --> {r['end']}", r["telugu"], r["english"], r.get("correction") or ""]
         for r in rows],
        REVIEW_WIDTHS,
        telugu_cols=["C"],
    )

    issues = [r for r in rows if (r.get("review") or {}).get("verdict") == "issue"]
    if any(r.get("review") for r in rows):
        _write_sheet(
            wb.create_sheet("AI Review"),
            ["Cue Number", "Severity", "Category", "Issue", "Reviewed English", "AI Suggested English"],
            [[r["cue"], r["review"]["severity"], r["review"]["category"], r["review"]["issue"],
              r["review"]["reviewed"], r["review"]["suggestion"]] for r in issues],
            {"A": 12, "B": 12, "C": 14, "D": 60, "E": 60, "F": 60},
        )
    wb.save(out_path)


def english_subtitles(rows: List[Dict[str, Any]], fmt: str) -> str:
    """Final English subtitles: the human correction wins over the AI draft."""
    return format_subtitles(((r["start"], r["end"], current_english(r)) for r in rows), fmt)


# --------------------------------------------------------------------------- #
# CLI convenience
# --------------------------------------------------------------------------- #
def translate_to_files(
    *,
    source_text: str,
    source_name: str,
    out_dir: Path,
    skill_text: str,
    model: Optional[str] = None,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    overlap: int = DEFAULT_OVERLAP,
    progress: ProgressFn = _noop,
) -> Dict[str, Path]:
    """Parse, translate, and write brief/workbook/raw JSON to out_dir.

    Returns a dict of artifact name -> path.
    """
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    cues = parse_source_text(source_text)
    result = translate_cues(
        cues,
        skill_text=skill_text,
        model=model,
        chunk_size=chunk_size,
        overlap=overlap,
        progress=progress,
    )

    stem = safe_stem(source_name)
    brief_path = out_dir / f"{stem}_discourse_brief.md"
    workbook_path = out_dir / f"{stem}_master_review.xlsx"
    raw_path = out_dir / f"{stem}_translation_raw.json"

    brief_path.write_text(result["brief"], encoding="utf-8")
    create_workbook(result["rows"], workbook_path)
    raw_path.write_text(
        json.dumps(result["all_items"], ensure_ascii=False, indent=2), encoding="utf-8"
    )

    progress("Done.")
    return {"brief": brief_path, "workbook": workbook_path, "raw": raw_path}
